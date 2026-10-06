from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document

import business_requirements_agent as brd
from project_records import load_project_records
from proposal_scope import preserve_source_bullets, requested_products


SAMPLE = (Path(__file__).parent / "fixtures/zoho_event_requirement.txt").read_text()


class BusinessRequirementsAgentTests(unittest.TestCase):
    def setUp(self):
        records, _ = load_project_records()
        products = requested_products(SAMPLE, records)
        self.sections = preserve_source_bullets([], SAMPLE, products)
        self.analysis = {
            "summary": "The business wants a connected system for its event work.",
            "requirement_sections": self.sections,
            "questions": [{"id": "q1", "question": "How should work be scheduled?"}],
            "commercial_categories": ["Zoho CRM", "Training"],
        }
        self.answers = {"q1": "Leave open for discovery"}

    def build(self, requirement=SAMPLE, analysis=None):
        with patch.object(brd, "_draft_brief",
                          side_effect=lambda sections, answers, fallback: {
                              "overview": fallback, "process_views": []}):
            return brd.build_document(requirement, analysis or self.analysis,
                                      self.answers, "Sample requirement.txt")

    def test_shared_product_heading_does_not_duplicate_the_same_checklist(self):
        doc = self.build()
        areas = [row["name"] for row in doc["requirements_by_area"]]
        self.assertIn("Zoho Marketing Automation & Zoho Campaigns", areas)
        self.assertIn("Zoho Creator & Zoho Backstage", areas)
        self.assertEqual(len(doc["requirements"]), 47)
        self.assertEqual(len({row["id"] for row in doc["requirements"]}), 47)

    def test_different_product_checklists_remain_separate(self):
        requirement = "Zoho Creator and Zoho Backstage are both requested."
        analysis = dict(self.analysis, requirement_sections=[
            {"product": "Zoho Creator", "requirements": ["Build custom forms"]},
            {"product": "Zoho Backstage", "requirements": ["Manage event pages"]},
        ])
        doc = self.build(requirement, analysis)
        names = [row["name"] for row in doc["requirements_by_area"]]
        self.assertIn("Zoho Creator", names)
        self.assertIn("Zoho Backstage", names)
        self.assertNotIn("Zoho Creator & Zoho Backstage", names)

    def test_word_export_contains_requirements_and_skips_empty_sections(self):
        doc = self.build()
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "brd.docx"
            brd.build_docx(doc, str(path))
            exported = Document(path)
        text = "\n".join(p.text for p in exported.paragraphs)
        text += "\n" + "\n".join(cell.text for table in exported.tables
                                   for row in table.rows for cell in row.cells)
        self.assertIn("Zoho Creator & Zoho Backstage", text)
        self.assertIn("Event-wise project setup", text)
        self.assertEqual(text.count("CRM Integration"), 1)
        self.assertNotIn("No specific requirements were stated", text)
        self.assertNotIn("Training requirements were not specified", text)

    def test_html_keeps_open_question_and_requested_costs(self):
        doc = self.build()
        html = brd.build_html(doc)
        self.assertIn("Questions and Decisions", html)
        self.assertIn("To be confirmed", html)
        self.assertIn("Commercial Items Requested", html)

    def test_acceptance_criteria_are_read_from_flattened_document_table(self):
        text = "Project heading\n9.2 Document Sign-Off\n# | Acceptance Criterion | Verified By | Status\n"
        text += "1 | Lead assignment works as agreed | Sales Manager | Pending\n"
        text += "Name & Designation | Organization | Signature | Date\n"
        self.assertEqual(brd._acceptance_table_rows(text), ["Lead assignment works as agreed"])
        doc = self.build(SAMPLE + "\n" + text)
        self.assertEqual(doc["acceptance_criteria"], ["Lead assignment works as agreed"])


if __name__ == "__main__":
    unittest.main()
