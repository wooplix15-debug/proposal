import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

import software_planning_agent as agent


class SoftwarePlanningAgentTests(unittest.TestCase):
    def test_analysis_keeps_only_verbatim_requirements_and_does_not_infer_title(self):
        source = "Build a repair booking app. Customers can upload photos and track repair status."
        with patch.object(agent, "_call", return_value={
            "summary": "A repair booking app.",
            "requirements": ["Customers can upload photos", "Invented billing portal"],
            "requirement_sections": [{"product": "Bookings", "requirements": ["Customers can upload photos"]}],
            "questions": [{"question": "Who handles bookings?", "reason": "This changes access.",
                           "options": ["Staff", "Customer self-service"]}],
        }):
            result = agent.analyze(source)
        self.assertEqual(result["project_name"], "")
        self.assertEqual([row["statement"] for row in result["requirements"]],
                         ["Build a repair booking app.", "Customers can upload photos and track repair status."])
        self.assertEqual(result["requirements"][0]["id"], agent._req_id("Build a repair booking app."))
        self.assertEqual(result["questions"][0]["id"], "q1")
        self.assertEqual(result["requirement_sections"][0]["product"], "Bookings")

    def test_explicit_title_and_user_title_are_respected(self):
        with patch.object(agent, "_call", return_value={"requirements": [], "questions": []}):
            from_source = agent.analyze("Project Name: Field Service Portal\nBuild a portal for customers.")
            supplied = agent.analyze("Build a portal for customers.", "Client Chosen Name")
        self.assertEqual(from_source["project_name"], "Field Service Portal")
        self.assertEqual(supplied["project_name"], "Client Chosen Name")

    def test_pack_has_eight_documents_and_uses_existing_proposal_and_brd_builders(self):
        analysis = {
            "project_name": "Repair Portal", "client_name": "Northstar",
            "summary": "Customers submit repair requests.", "questions": [],
            "requirements": [
                {"id": "REQ-A", "statement": "Customers submit repair requests.", "source": "client brief"},
                {"id": "REQ-B", "statement": "Staff assign each request to a technician.", "source": "client brief"},
            ],
        }
        model_result = {
            "ux": {"principles": ["Keep request status visible."]},
            "implementation": {"architecture_summary": "Decisions remain open.",
                               "backlog": [{"title": "Request intake", "requirement_ids": ["REQ-A"]}]},
            "qa": {"test_cases": [{"scenario": "Submit a request", "steps": ["Submit a request"],
                                    "requirement_ids": ["REQ-A"]}]},
            "training": {"training_topics": ["Submitting repair requests"]},
            "documentation": {"training_topics": ["Submitting repair requests"]},
            "handover": {"next_steps": ["Review scope."]},
        }
        proposal = {"client": {"company_name": "Northstar", "project_name": "Repair Portal"},
                    "scope": [], "project_introduction": "Repair portal."}
        brd = {"title": "Business Requirements Document", "project_name": "Repair Portal",
               "status": "Draft for business review", "prepared_by": "Wooplix", "date": "today",
               "summary": "Repair portal", "requirements": [], "requirements_by_area": [],
               "process_views": [], "specification_tables": [], "current_state": [], "objectives": [],
               "stakeholders": [], "scope_items": [], "purpose": [], "acceptance_criteria": [],
               "training_requirements": [], "open_decisions": []}
        with patch.object(agent, "_call", return_value=model_result), \
             patch.object(agent.brand, "draft_proposal", return_value=proposal) as proposal_builder, \
             patch("business_requirements_agent.build_document", return_value=brd) as brd_builder:
            docs = agent.build_pack(analysis)
        self.assertEqual([doc["kind"] for doc in docs],
                         ["requirements", "proposal", "brd", "ux", "implementation", "qa", "documentation", "handover"])
        proposal_builder.assert_called_once()
        brd_builder.assert_called_once()
        self.assertEqual(docs[1]["content"]["client"]["project_name"], "Repair Portal")
        self.assertEqual(docs[2]["content"]["title"], brd["title"])
        dev_ids = {rid for row in docs[4]["content"]["backlog"] for rid in row["requirement_ids"]}
        test_ids = {rid for row in docs[5]["content"]["test_cases"] for rid in row["requirement_ids"]}
        self.assertEqual(dev_ids, {"REQ-A", "REQ-B"})
        self.assertEqual(test_ids, {"REQ-A", "REQ-B"})
        self.assertEqual(docs[5]["content"]["test_cases"][-1]["status"], "Planned — not executed")
        self.assertEqual(docs[-1]["content"]["implementation_status"], "Not built or deployed")

    def test_answer_validation_requires_a_response_or_leave_open_for_each_question(self):
        analysis = {"questions": [{"id": "q1", "question": "Who uses it?"}], "requirements": []}
        with self.assertRaisesRegex(ValueError, "Answer each question"):
            agent.build_pack(analysis, {})

    def test_docx_and_html_render_a_clear_status_and_escape_source_text(self):
        doc = {"kind": "requirements", "title": "Requirements & Scope", "project_name": "",
               "client_name": "Client", "status": "Planning draft — for review", "prepared_by": "Wooplix",
               "date": "13 October 2026", "content": {"requirements": [{"id": "REQ-A", "statement": "<script>unsafe</script>"}]}}
        rendered = agent.build_html(doc)
        self.assertNotIn("<script>unsafe", rendered)
        self.assertIn("Planning draft", rendered)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "requirements.docx"
            agent.build_docx(doc, str(path))
            exported = Document(path)
        text = "\n".join(p.text for p in exported.paragraphs)
        self.assertIn("Requirements", text)
        self.assertIn("not provided", "\n".join(c.text for t in exported.tables for r in t.rows for c in r.cells).lower())


if __name__ == "__main__":
    unittest.main()
