"""Independent agent for creating a source-grounded Business Requirements Document."""
from __future__ import annotations

import html
import json
import os
import re
import shutil
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import wooplix_agent as brand


def _section_items(sections):
    """Assign stable BRD IDs to every distinct source requirement."""
    rows = []
    seen = set()
    number = 1
    for section in sections or []:
        product = str(section.get("product") or "Business requirement").strip()
        for requirement in section.get("requirements") or []:
            text = re.sub(r"\s+", " ", str(requirement)).strip()
            key = (product.casefold(), text.casefold())
            if not text or key in seen:
                continue
            seen.add(key)
            rows.append({"id": f"BR-{number:03}", "area": product, "requirement": text})
            number += 1
    return rows


def _source_list(requirement, section, next_headers):
    """Read literal list items beneath one source heading."""
    lines = requirement.splitlines()
    start = next((i for i, line in enumerate(lines)
                  if re.match(rf"^\s*(?:\d+[.)]\s*)?{section}\b", line, re.I)), None)
    if start is None:
        return []
    result = []
    for line in lines[start + 1:]:
        if (re.match(r"^\s*\d+[.)]\s+", line)
                or any(re.match(rf"^\s*{header}\b", line, re.I) for header in next_headers)):
            break
        text = re.sub(r"^\s*(?:[-•*]|\d+[.)])\s*", "", line).strip()
        if text:
            result.append(text)
    return list(dict.fromkeys(result))


def _draft_brief(sections, answers, fallback_summary):
    """Dedicated BRD-agent pass: explain groupings without creating new scope."""
    try:
        from groq import Groq, RateLimitError

        system = """You are Wooplix's Business Requirements Document agent. This is a
separate requirements-analysis workflow, not the proposal writer. Use only the
reviewed requirement checklist and user answers supplied below. Do not use external
product knowledge. Do not add features, software products, integrations, roles,
workflows, rules, metrics or commitments. Preserve uncertainty as uncertainty.
Return JSON with one concise factual overview and process_views. Each process_view
must use an exact existing business-area heading and include only a short explanation
of activities explicitly listed for that area. Add requirement_ids from the exact
IDs supplied. Cross-application relationships can be described only when an explicit
requirement states them. If the source does not define a sequence, do not make one.
Schema: {"overview":"", "process_views":[{"area":"","description":"",
"requirement_ids":["BR-001"]}]}"""
        user = json.dumps({"requirements": [
            {"id": item["id"], "area": item["area"], "statement": item["requirement"]}
            for item in _section_items(sections)],
            "answers": answers}, ensure_ascii=False)
        models = list(dict.fromkeys([brand.GROQ_MODEL, "openai/gpt-oss-120b",
                                     "openai/gpt-oss-20b", "qwen/qwen3.8-27b"]))
        keys = brand.GROQ_API_KEYS or ([brand.GROQ_API_KEY] if brand.GROQ_API_KEY else [])
        for model in models:
            for key in keys:
                try:
                    kwargs = {"model": model,
                              "messages": [{"role": "system", "content": system},
                                           {"role": "user", "content": user}],
                              "temperature": 0.1, "max_completion_tokens": 2200,
                              "response_format": {"type": "json_object"}}
                    if "gpt-oss" in model:
                        kwargs["reasoning_effort"] = "low"
                    response = Groq(api_key=key, timeout=45, max_retries=0).chat.completions.create(**kwargs)
                    text = response.choices[0].message.content or "{}"
                    data = json.loads(text)
                    by_area = {}
                    for item in _section_items(sections):
                        by_area.setdefault(item["area"], set()).add(item["id"])
                    views = []
                    for view in data.get("process_views", []):
                        area = str(view.get("area", "")).strip()
                        ids = view.get("requirement_ids", [])
                        if (area in by_area and isinstance(ids, list) and ids
                                and set(ids) <= by_area[area]):
                            views.append({"area": area,
                                          "description": str(view.get("description", ""))[:700],
                                          "requirement_ids": ids})
                    return {"overview": str(data.get("overview") or fallback_summary)[:1600],
                            "process_views": views}
                except RateLimitError:
                    break
                except Exception:
                    continue
    except Exception:
        pass
    return {"overview": fallback_summary, "process_views": []}


def build_document(requirement, analysis, answers, source_name=""):
    """Make a detailed BRD from the reviewed checklist without expanding scope."""
    questions = analysis.get("questions") or []
    if set(answers) != {q["id"] for q in questions}:
        raise ValueError("Answer each question or choose Leave open.")
    for answer in answers.values():
        if not isinstance(answer, str) or not answer.strip() or len(answer) > 4000:
            raise ValueError("Enter each answer (up to 4000 characters).")

    sections = analysis.get("requirement_sections") or []
    requirements = _section_items(sections)
    brief = _draft_brief(sections, answers, str(analysis.get("summary") or "").strip())
    activities = []
    raw_text = requirement.casefold()
    shared_pairs = [
        (("Zoho Marketing Automation", "Zoho Campaigns"),
         bool(re.search(r"zoho\s+marketing\s+automation\s*[,/&]\s*campaign", raw_text))),
        (("Zoho Creator", "Zoho Backstage"),
         bool(re.search(r"zoho\s+creator\s*[,/&]\s*backstage", raw_text))),
    ]
    consumed = set()
    for section in sections:
        product = str(section.get("product") or "").strip()
        items = list(dict.fromkeys(str(x).strip() for x in section.get("requirements", []) if str(x).strip()))
        if product in consumed:
            continue
        for pair, explicitly_shared in shared_pairs:
            if explicitly_shared and product in pair:
                peers = [x for x in sections if x.get("product") in pair]
                shared = list(dict.fromkeys(
                    str(item).strip() for peer in peers for item in peer.get("requirements", [])
                    if str(item).strip()))
                if len(peers) > 1 and all(
                    list(dict.fromkeys(str(i).strip() for i in peer.get("requirements", []) if str(i).strip())) == shared
                    for peer in peers
                ):
                    product = " & ".join(pair)
                    items = shared
                    consumed.update(pair)
                    break
        if items:
            activities.append({"area": product, "items": items})

    def matches(*terms):
        return [{"area": row["area"], "items": [item for item in row["items"]
                if any(term in item.casefold() for term in terms)]}
                for row in activities if any(any(term in item.casefold() for term in terms)
                                               for item in row["items"])]

    # Keep stated business objectives verbatim when the source contains that section.
    objectives = _source_list(requirement, "Main Objective", ["Cost Proposal", "Zoho CRM", "Training"])
    # Use the reviewed, source-filtered checklist as the authority for every product.
    areas = [{"name": row["area"], "requirements": row["items"]} for row in activities]
    open_decisions = []
    for question in questions:
        answer = answers[question["id"]].strip()
        if answer.casefold() in {"leave open for discovery", "confirm during discovery", ""}:
            open_decisions.append({"question": question["question"], "answer": "To be confirmed"})
        else:
            open_decisions.append({"question": question["question"], "answer": answer})

    return {
        "title": "Business Requirements Document",
        "project_name": (Path(source_name).stem.replace("_", " ").strip()
                         if source_name and Path(source_name).stem.casefold() not in {"pasted requirement", "requirement"}
                         else "Project name to be confirmed"),
        "prepared_for": "",
        "prepared_by": brand.COMPANY_NAME,
        "date": brand._ordinal_day(datetime.now(ZoneInfo("Asia/Kolkata"))),
        "status": "Draft for business review",
        "summary": brief["overview"],
        "process_views": brief["process_views"],
        "objectives": objectives,
        "areas": areas,
        "requirements": requirements,
        "business_rules": [row for row in requirements
                           if any(term in row["requirement"].casefold()
                                  for term in ("duplicate", "same company", "billable / non-billable", "planned vs actual"))],
        "data_requirements": matches("data", "record", "duplicate", "contact", "lead", "company", "event-wise", "batch-wise", "timesheet", "work notes"),
        "integration_requirements": matches("integration", "integrat", "whatsapp", "sms", "crm integration"),
        "reporting_requirements": matches("report", "dashboard", "analytics", "tracking"),
        "access_requirements": matches("role", "permission", "assignment", "assigned"),
        "training_requirements": next((x["items"] for x in activities if x["area"].casefold() == "training"), []),
        "commercial_categories": analysis.get("commercial_categories") or [],
        "open_decisions": open_decisions,
        "source_name": Path(source_name).name if source_name else "",
        "activity_map": activities,
    }


def _e(value):
    return html.escape(str(value or ""), quote=True)


def build_html(doc):
    """Render a separate BRD in Wooplix's branded document style."""
    main = brand._logo_data_uri()
    badge = brand._badge_data_uri()
    main_img = f'<img src="{main}" style="height:38px" alt="Wooplix">' if main else ""
    badge_img = f'<img src="{badge}" style="height:26px" alt="Partner">' if badge else ""
    out = [f"""<!doctype html><html><head><meta charset="utf-8"><style>
@page {{ size:A4; margin:27mm 15mm 22mm; }}
body {{ font-family:'DejaVu Sans',sans-serif; font-size:9pt; line-height:1.45; color:#1e293b; }}
.header {{ position:fixed; top:-20mm; width:100%; }} .header td:last-child {{text-align:right}}
.footer {{position:fixed;bottom:-8mm;width:100%;border-top:1px solid #cbd5e1;font-size:7pt;color:#64748b}}
.footer td {{padding-top:3mm}} h1 {{font-size:19pt;color:#002b49;margin:12mm 0 8mm}}
h2 {{page-break-after:avoid;background:#1a365d;color:#fff;font-size:11pt;padding:7px 10px;margin:18px 0 8px}}
h3 {{page-break-after:avoid;color:#008080;font-size:10pt;border-bottom:1px solid #008080;padding-bottom:3px;margin:13px 0 5px}}
p {{margin:4px 0 8px}} table {{width:100%;border-collapse:collapse;margin:6px 0 12px;font-size:8.5pt}}
thead {{display:table-header-group}} tr {{page-break-inside:avoid}} th {{background:#1a365d;color:white;text-align:left;font-size:8pt;padding:6px;border:1px solid #1a365d}}
td {{padding:5px 7px;vertical-align:top;border:1px solid #cbd5e1}} tbody tr:nth-child(even) td {{background:#f8fafc}}
.meta td:first-child {{background:#edf3f8;font-weight:bold;width:30%}} .muted {{color:#64748b;font-size:8pt}}
</style></head><body><table class="header"><tr><td>{main_img}</td><td>{badge_img}</td></tr></table>
<table class="footer"><tr><td><b>{_e(brand.COMPANY_NAME)}</b> · Confidential</td><td align="center">{_e(brand.COMPANY_EMAIL)}</td><td align="right">{_e(brand.COMPANY_WEBSITE)}</td></tr></table>
<h1>{_e(doc['title'])}</h1>
<table class="meta"><tbody>"""]
    for label, value in (("Project", doc["project_name"]), ("Status", doc["status"]),
                         ("Prepared by", doc["prepared_by"]), ("Date", doc["date"])):
        out.append(f"<tr><td>{_e(label)}</td><td>{_e(value)}</td></tr>")
    out.append("</tbody></table>")

    def section(title):
        out.append(f"<h2>{_e(title)}</h2>")

    def list_table(rows, first="Requirement"):
        out.append(f"<table><thead><tr><th style='width:12%'>ID</th><th style='width:24%'>Business area</th><th>{_e(first)}</th></tr></thead><tbody>")
        for row in rows:
            out.append(f"<tr><td>{_e(row['id'])}</td><td>{_e(row['area'])}</td><td>{_e(row['requirement'])}</td></tr>")
        out.append("</tbody></table>")

    section("1. Project Overview")
    out.append(f"<p>{_e(doc['summary'])}</p>")
    section("2. Business Objectives")
    out.append("<ul>")
    for item in doc["objectives"]:
        out.append(f"<li>{_e(item)}</li>")
    if not doc["objectives"]:
        out.append("<li>The overall business objectives require confirmation by the business owner.</li>")
    out.append("</ul>")
    section("3. Scope and Business Areas")
    for area in doc["areas"]:
        out.append(f"<h3>{_e(area['name'])}</h3><ul>")
        for item in area["requirements"]:
            out.append(f"<li>{_e(item)}</li>")
        out.append("</ul>")
    if doc["process_views"]:
        section("4. Business Process View (Draft)")
        out.append("<p>These short descriptions organize the stated work areas. They are not a new scope or an assumed sequence; confirm the interpretation during business review.</p>")
        for view in doc["process_views"]:
            out.append(f"<h3>{_e(view['area'])}</h3><p>{_e(view['description'])}</p>")
            out.append(f"<p class='muted'>Requirements: {_e(', '.join(view['requirement_ids']))}</p>")
    section("5. Detailed Functional Requirements")
    out.append("<p>Requirement statements below preserve the reviewed request. Each BRD ID can be used to discuss scope and record approval.</p>")
    list_table(doc["requirements"])

    section("6. Business Rules")
    if doc["business_rules"]:
        list_table(doc["business_rules"])
    else:
        out.append("<p class='muted'>No specific business rules were stated in the request.</p>")
    for title, key in (("7. Data and Information Requirements", "data_requirements"),
                       ("8. Integrations", "integration_requirements"),
                       ("9. Reporting and Analytics", "reporting_requirements"),
                       ("10. User Roles and Access", "access_requirements")):
        section(title)
        rows = [item for group in doc[key] for item in group["items"]]
        if rows:
            out.append("<ul>" + "".join(f"<li>{_e(item)}</li>" for item in rows) + "</ul>")
        else:
            out.append("<p class='muted'>No specific requirements were stated in the request.</p>")
    section("11. Training Requirements")
    if doc["training_requirements"]:
        out.append("<ul>" + "".join(f"<li>{_e(x)}</li>" for x in doc["training_requirements"]) + "</ul>")
    else:
        out.append("<p class='muted'>Training requirements were not specified.</p>")
    section("12. Commercial Items Requested")
    out.append("<p>Amounts were not included in the requirements. The requested quote categories are:</p><ul>")
    out.extend(f"<li>{_e(x)}</li>" for x in doc["commercial_categories"])
    out.append("</ul>")
    section("13. Decisions and Details to Confirm")
    out.append("<table><thead><tr><th>Open item</th><th>Current status / answer</th></tr></thead><tbody>")
    if doc["open_decisions"]:
        for row in doc["open_decisions"]:
            out.append(f"<tr><td>{_e(row['question'])}</td><td>{_e(row['answer'])}</td></tr>")
    else:
        out.append("<tr><td colspan='2'>No open items were recorded during this review.</td></tr>")
    out.append("</tbody></table>")
    section("14. Review and Sign-off")
    out.append("<p>Review each BRD requirement with the business owner. Confirm the workflow details, data rules, access, integrations, reports, and any open items before implementation scope is approved.</p>")
    out.append("<table><tbody><tr><td style='width:50%;height:48px'>Business owner / date</td><td>Wooplix / date</td></tr></tbody></table>")
    out.append("</body></html>")
    return "".join(out)


def build_docx(doc, output):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    report = Document()
    section = report.sections[0]
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.left_margin = section.right_margin = Inches(0.7)
    section.top_margin = section.bottom_margin = Inches(0.75)
    normal = report.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)
    header = section.header
    logo_table = header.add_table(rows=1, cols=2, width=Inches(7.1))
    if Path(brand.LOGO_MAIN_PATH).exists():
        logo_table.cell(0, 0).paragraphs[0].add_run().add_picture(brand.LOGO_MAIN_PATH, width=Inches(2.0))
    if Path(brand.LOGO_BADGE_PATH).exists():
        p_badge = logo_table.cell(0, 1).paragraphs[0]
        p_badge.alignment = 2
        p_badge.add_run().add_picture(brand.LOGO_BADGE_PATH, width=Inches(1.65))
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    footer.add_run(f"{brand.COMPANY_NAME}  ·  Confidential  ·  {brand.COMPANY_WEBSITE}").font.size = Pt(8)
    p = report.add_paragraph()
    p.alignment = 1
    r = p.add_run(doc["title"])
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(0, 43, 73)
    table = report.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in (("Project", doc["project_name"]), ("Status", doc["status"]),
                         ("Prepared by", doc["prepared_by"]), ("Date", doc["date"])):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, str(value)
        cells[0]._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="edf3f8"/>'))

    def heading(value):
        h = report.add_heading(value, level=1)
        for run in h.runs:
            run.font.color.rgb = RGBColor(0x1a, 0x36, 0x5d)
        return h

    def bullets(values):
        for value in values:
            report.add_paragraph(str(value), style="List Bullet")

    heading("1. Project Overview")
    report.add_paragraph(doc["summary"])
    heading("2. Business Objectives")
    bullets(doc["objectives"] or ["The overall business objectives require confirmation by the business owner."])
    heading("3. Scope and Business Areas")
    for area in doc["areas"]:
        report.add_heading(area["name"], level=2)
        bullets(area["requirements"])
    if doc["process_views"]:
        heading("4. Business Process View (Draft)")
        report.add_paragraph("These descriptions organize the stated work areas. They are not new scope or an assumed sequence; confirm the interpretation during business review.")
        for view in doc["process_views"]:
            report.add_heading(view["area"], level=2)
            report.add_paragraph(view["description"])
            report.add_paragraph("Requirements: " + ", ".join(view["requirement_ids"]))
    heading("5. Detailed Functional Requirements")
    report.add_paragraph("Requirement statements below preserve the reviewed request. Each BRD ID can be used to discuss scope and record approval.")
    table = report.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, ("ID", "Business area", "Requirement")):
        cell.text = label
    for row in doc["requirements"]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = row["id"], row["area"], row["requirement"]
    heading("6. Business Rules")
    if doc["business_rules"]:
        table = report.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for cell, label in zip(table.rows[0].cells, ("ID", "Business area", "Rule")):
            cell.text = label
        for row in doc["business_rules"]:
            cells = table.add_row().cells
            cells[0].text, cells[1].text, cells[2].text = row["id"], row["area"], row["requirement"]
    else:
        report.add_paragraph("No specific business rules were stated in the request.")
    for title, key in (("7. Data and Information Requirements", "data_requirements"),
                       ("8. Integrations", "integration_requirements"),
                       ("9. Reporting and Analytics", "reporting_requirements"),
                       ("10. User Roles and Access", "access_requirements")):
        heading(title)
        items = [item for group in doc[key] for item in group["items"]]
        bullets(items or ["No specific requirements were stated in the request."])
    heading("11. Training Requirements")
    bullets(doc["training_requirements"] or ["Training requirements were not specified."])
    heading("12. Commercial Items Requested")
    report.add_paragraph("Amounts were not included in the requirements. The requested quote categories are:")
    bullets(doc["commercial_categories"])
    heading("13. Decisions and Details to Confirm")
    table = report.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Open item", "Current status / answer"
    for row in doc["open_decisions"]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = row["question"], row["answer"]
    heading("14. Review and Sign-off")
    report.add_paragraph("Review each BRD requirement with the business owner. Confirm workflow details, data rules, access, integrations, reports, and open items before implementation scope is approved.")
    signoff = report.add_table(rows=1, cols=2)
    signoff.style = "Table Grid"
    signoff.rows[0].cells[0].text = "Business owner / date"
    signoff.rows[0].cells[1].text = "Wooplix / date"
    report.save(output)
    return output


def build_pdf(doc, output):
    php = shutil.which("php")
    script = Path(brand.PHP_SCRIPT)
    if not php or not (Path(brand.HERE) / "vendor" / "autoload.php").exists():
        raise RuntimeError("The PDF renderer is unavailable.")
    fd, html_path = tempfile.mkstemp(suffix=".html", prefix="wooplix_brd_", dir=str(Path(output).parent))
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            stream.write(build_html(doc))
        result = subprocess.run([php, str(script), html_path, str(output)],
                                capture_output=True, text=True, timeout=240)
        if result.returncode or not Path(output).exists():
            raise RuntimeError("Could not render the Business Requirements Document PDF.")
    finally:
        try:
            os.unlink(html_path)
        except OSError:
            pass
    return output
